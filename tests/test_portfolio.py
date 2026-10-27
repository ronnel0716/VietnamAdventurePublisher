import tempfile
import unittest
from pathlib import Path
from docx import Document
from vap.portfolio import extract, render, main


class PipelineTests(unittest.TestCase):
    def test_extract_preserves_mixed_order_and_unicode(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'fixture.docx'
            doc = Document()
            doc.add_heading('Café ₱', 0)
            doc.add_heading('Plan', 1)
            doc.add_table(rows=1, cols=1).cell(0, 0).text = '₱250'
            doc.add_paragraph('After table')
            doc.save(path)
            data = extract(path)
            self.assertEqual(data['title'], 'Café ₱')
            self.assertEqual([b['type'] for b in data['blocks']], ['paragraph', 'heading', 'table', 'paragraph'])
            self.assertEqual(data['blocks'][2]['rows'][0][0], '₱250')
            main([str(path), '--output-dir', str(Path(directory) / 'out')])
            with self.assertRaises(SystemExit):
                main([str(path), '--output-dir', str(Path(directory) / 'out')])

    def test_html_escapes_document_content(self):
        result = render({'title':'<script>', 'blocks':[{'type':'paragraph','text':'<img onerror="bad">'}]})
        self.assertNotIn('<script>', result)
        self.assertIn('&lt;img', result)

    def test_rejects_invalid_input(self):
        with self.assertRaises(SystemExit):
            main(['missing.docx'])


if __name__ == '__main__':
    unittest.main()
