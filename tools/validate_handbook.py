
r"""Validate a Vietnam Adventure handbook before publication.

Usage:
    .\.venv\Scripts\python.exe tools\validate_handbook.py
    .\.venv\Scripts\python.exe tools\validate_handbook.py path\to\handbook.docx
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter
from dataclasses import dataclass
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
SOURCE_FOLDER = ROOT / "handbook" / "source"

CHAPTER_PATTERN = re.compile(r"^CHAPTER\s+(\d+)\s*$", re.IGNORECASE)
PLACEHOLDER_PATTERN = re.compile(
    r"\[[^\]]*(?:placeholder|todo|tbd|tbc)[^\]]*\]",
    re.IGNORECASE,
)


@dataclass(frozen=True, slots=True)
class ValidationIssue:
    severity: str
    message: str
    paragraph_index: int | None = None


def find_source() -> Path:
    """Return the only DOCX inside handbook/source, if unambiguous."""
    candidates = sorted(SOURCE_FOLDER.glob("*.docx"))

    if not candidates:
        raise FileNotFoundError(
            f"No DOCX handbook found in: {SOURCE_FOLDER}"
        )

    if len(candidates) > 1:
        names = "\n".join(f"  - {item.name}" for item in candidates)
        raise ValueError(
            "Multiple DOCX files found. Provide the source path explicitly:\n"
            f"{names}"
        )

    return candidates[0]


def is_placeholder(text: str) -> bool:
    return bool(PLACEHOLDER_PATTERN.search(text))


def find_chapter_title(paragraphs, chapter_index: int) -> str | None:
    """Find the first useful heading after a chapter marker."""
    for paragraph in paragraphs[chapter_index + 1 :]:
        text = paragraph.text.strip()

        if CHAPTER_PATTERN.fullmatch(text):
            return None

        if not text or is_placeholder(text):
            continue

        return text

    return None


def validate(source: Path) -> tuple[Document, list[ValidationIssue], list[int]]:
    """Inspect a source handbook and return its issues and chapter positions."""
    if not source.exists():
        raise FileNotFoundError(f"Source handbook not found: {source}")

    document = Document(source)
    paragraphs = document.paragraphs
    issues: list[ValidationIssue] = []
    chapter_indexes: list[int] = []

    title = next((p.text.strip() for p in paragraphs if p.text.strip()), "")

    if not title:
        issues.append(ValidationIssue("ERROR", "Handbook title is missing."))

    for index, paragraph in enumerate(paragraphs):
        text = paragraph.text.strip()

        if is_placeholder(text):
            issues.append(
                ValidationIssue(
                    "ERROR",
                    "Unresolved placeholder: " + text,
                    index,
                )
            )

        if CHAPTER_PATTERN.fullmatch(text):
            chapter_indexes.append(index)

    if not chapter_indexes:
        issues.append(ValidationIssue("ERROR", "No chapters detected."))

    expected_number = 1

    for chapter_index in chapter_indexes:
        marker = paragraphs[chapter_index].text.strip()
        number = int(CHAPTER_PATTERN.fullmatch(marker).group(1))

        if number != expected_number:
            issues.append(
                ValidationIssue(
                    "ERROR",
                    f"Expected Chapter {expected_number}, found Chapter {number}.",
                    chapter_index,
                )
            )

        expected_number = number + 1

        if find_chapter_title(paragraphs, chapter_index) is None:
            issues.append(
                ValidationIssue(
                    "ERROR",
                    f"Chapter {number} has no usable title.",
                    chapter_index,
                )
            )

    return document, issues, chapter_indexes


def print_issue_group(title: str, issues: list[ValidationIssue]) -> None:
    print()
    print(title)
    print("-" * len(title))

    if not issues:
        print("None")
        return

    for issue in issues:
        location = ""

        if issue.paragraph_index is not None:
            location = f" [paragraph {issue.paragraph_index}]"

        print(f"- {issue.message}{location}")


def print_placeholder_summary(
    issues: list[ValidationIssue],
    show_details: bool,
) -> None:
    print()
    print("Unresolved Placeholders")
    print("-----------------------")

    if not issues:
        print("None")
        return

    labels = [
        issue.message.removeprefix("Unresolved placeholder: ")
        for issue in issues
    ]

    for label, count in Counter(labels).most_common():
        print(f"- {count:>3}  {label}")

    if show_details:
        print("\nPlaceholder locations")
        print("---------------------")

        for issue in issues:
            print(f"- {issue.message} [paragraph {issue.paragraph_index}]")


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Validate a handbook source DOCX before publishing."
    )
    parser.add_argument(
        "source",
        nargs="?",
        type=Path,
        help="Optional source handbook DOCX. Defaults to handbook/source.",
    )
    parser.add_argument(
        "--details",
        action="store_true",
        help="Show the paragraph number for every unresolved placeholder.",
    )
    args = parser.parse_args()

    try:
        source = args.source.resolve() if args.source else find_source()
        document, issues, chapter_indexes = validate(source)
    except (FileNotFoundError, ValueError) as error:
        print(f"Validation could not start: {error}")
        return 2

    errors = [issue for issue in issues if issue.severity == "ERROR"]
    warnings = [issue for issue in issues if issue.severity == "WARNING"]
    placeholder_count = sum(
        1
        for issue in errors
        if issue.message.startswith("Unresolved placeholder:")
    )
    placeholder_issues = [
        issue
        for issue in errors
        if issue.message.startswith("Unresolved placeholder:")
    ]
    structural_errors = [
        issue
        for issue in errors
        if not issue.message.startswith("Unresolved placeholder:")
    ]

    print("=" * 60)
    print("Vietnam Adventure Publisher")
    print("Handbook Validation Report")
    print("=" * 60)
    print(f"\nSource     : {source}")
    print(f"Paragraphs : {len(document.paragraphs)}")
    print(f"Tables     : {len(document.tables)}")
    print(f"Images     : {len(document.inline_shapes)}")
    print(f"Chapters   : {len(chapter_indexes)}")
    print(f"Placeholders : {placeholder_count}")

    print_issue_group("Structural Errors", structural_errors)
    print_placeholder_summary(placeholder_issues, args.details)
    print_issue_group("Warnings", warnings)

    print("\nSummary")
    print("-------")

    if errors:
        print("STATUS : NOT READY")
        print("Resolve all errors before publishing the final handbook.")
        return 1

    if warnings:
        print("STATUS : VALID WITH WARNINGS")
        return 0

    print("STATUS : VALID")
    print("The handbook is ready for publishing.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
