r"""Create an AI-ready asset manifest from handbook placeholders.

Usage:
    .\.venv\Scripts\python.exe tools\generate_asset_manifest.py
    .\.venv\Scripts\python.exe tools\generate_asset_manifest.py path\to\handbook.docx
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parent.parent
SOURCE_FOLDER = ROOT / "handbook" / "source"
OUTPUT_FOLDER = ROOT / "output"

CHAPTER_PATTERN = re.compile(r"^CHAPTER\s+(\d+)\s*$", re.IGNORECASE)
PLACEHOLDER_PATTERN = re.compile(
    r"\[[^\]]*(?:placeholder|todo|tbd|tbc)[^\]]*\]",
    re.IGNORECASE,
)


@dataclass(slots=True)
class AssetTask:
    id: str
    status: str
    asset_type: str
    chapter_number: int | None
    chapter_title: str
    paragraph_index: int
    placeholder: str
    nearby_context: str
    prompt: str
    manual_input_required: bool
    notes: str


def find_source() -> Path:
    candidates = sorted(SOURCE_FOLDER.glob("*.docx"))

    if not candidates:
        raise FileNotFoundError(f"No DOCX found in: {SOURCE_FOLDER}")

    if len(candidates) > 1:
        names = ", ".join(item.name for item in candidates)
        raise ValueError(f"Multiple DOCX files found: {names}")

    return candidates[0]


def is_placeholder(text: str) -> bool:
    return bool(PLACEHOLDER_PATTERN.search(text))


def classify_asset(placeholder: str) -> tuple[str, bool, str]:
    normalized = placeholder.casefold()

    if "hero image" in normalized:
        return "hero_image", False, "Create as a 16:9 editorial travel image."
    if "navigation panel" in normalized:
        return "navigation_panel", False, "Create as a structured Word callout panel."
    if "icon set" in normalized:
        return "icon_set", False, "Create as a consistent monochrome icon set."
    if "timeline" in normalized:
        return "timeline_graphic", False, "Create as a clean horizontal itinerary timeline."
    if "packing planner" in normalized:
        return "packing_planner", False, "Create as a printable checklist worksheet."
    if "travel planner" in normalized or "budget planner" in normalized:
        return "travel_budget_planner", False, "Create as a printable planning worksheet."
    if "map" in normalized or "route" in normalized:
        return "route_map", False, "Create as a simplified route diagram, not a navigation map."
    if "reflection journal" in normalized:
        return "reflection_journal", False, "Create as a spacious guided reflection page."
    if "qr code" in normalized:
        return "qr_code", True, "A verified destination URL is required before creating a QR code."
    if "author:" in normalized or "publisher:" in normalized:
        return "publication_credits", True, "Requires final author, publisher, and design-credit details."
    if "handbook version" in normalized:
        return "publication_metadata", True, "Requires final version, publication date, and update URL."

    return "content_placeholder", True, "Review the source content and define the required replacement."


def chapter_context(paragraphs, index: int) -> tuple[int | None, str]:
    chapter_number: int | None = None
    chapter_title = "Front matter or end matter"

    for current_index in range(index, -1, -1):
        marker = paragraphs[current_index].text.strip()
        match = CHAPTER_PATTERN.fullmatch(marker)

        if not match:
            continue

        chapter_number = int(match.group(1))

        for candidate in paragraphs[current_index + 1 :]:
            text = candidate.text.strip()

            if CHAPTER_PATTERN.fullmatch(text):
                break

            if text and not is_placeholder(text):
                chapter_title = text
                break

        return chapter_number, chapter_title

    return chapter_number, chapter_title


def nearby_context(paragraphs, index: int) -> str:
    context: list[str] = []

    for current_index in range(max(0, index - 3), min(len(paragraphs), index + 4)):
        if current_index == index:
            continue

        text = paragraphs[current_index].text.strip()

        if text and not is_placeholder(text):
            context.append(text)

    return " | ".join(context[:3]) or "No nearby descriptive content found."


def build_prompt(
    asset_type: str,
    chapter_number: int | None,
    chapter_title: str,
    context: str,
    notes: str,
) -> str:
    chapter_label = (
        f"Chapter {chapter_number}: {chapter_title}"
        if chapter_number is not None
        else chapter_title
    )

    base = (
        "Vietnam Adventure 2026 travel handbook. "
        f"Context: {chapter_label}. "
        f"Nearby handbook content: {context}. "
    )

    prompts = {
        "hero_image": (
            "Create a premium documentary travel-photography hero image, "
            "landscape 16:9, authentic Vietnam setting relevant to the context, "
            "warm natural light, sophisticated editorial composition, no text, "
            "no logos, no watermarks, no collage."
        ),
        "navigation_panel": (
            "Create concise content for a polished navigation panel with "
            "three labeled areas: Chapter Overview, Key Stops, and QR Destination. "
            "Use only facts established in the nearby content; mark uncertain facts "
            "as review-needed."
        ),
        "icon_set": (
            "Create a cohesive set of six simple line icons for this chapter. "
            "Use one-color teal and navy styling, no text, clear at small print size."
        ),
        "timeline_graphic": (
            "Create a clean horizontal travel timeline showing the sequence of events "
            "already described in the nearby content. Use short labels and restrained colors."
        ),
        "packing_planner": (
            "Create a printable packing checklist organized by essentials, clothing, "
            "health, technology, and documents. Keep all entries relevant to this chapter."
        ),
        "travel_budget_planner": (
            "Create a printable trip-and-budget planning worksheet with realistic empty "
            "fields for dates, transport, accommodation, meals, activities, and contingency."
        ),
        "route_map": (
            "Create a simplified editorial route diagram showing only the locations and "
            "movement stated in the nearby content. Do not invent geography or routes."
        ),
        "reflection_journal": (
            "Create a guided reflection page with a short prompt, generous writing space, "
            "and a restrained travel-handbook visual style."
        ),
    }

    instruction = prompts.get(asset_type, notes)
    return base + instruction


def build_manifest(source: Path) -> list[AssetTask]:
    document = Document(source)
    paragraphs = document.paragraphs
    tasks: list[AssetTask] = []

    for index, paragraph in enumerate(paragraphs):
        placeholder = paragraph.text.strip()

        if not is_placeholder(placeholder):
            continue

        asset_type, manual_input_required, notes = classify_asset(placeholder)
        chapter_number, chapter_title = chapter_context(paragraphs, index)
        context = nearby_context(paragraphs, index)

        task = AssetTask(
            id=f"ASSET-{len(tasks) + 1:03d}",
            status="needs_input" if manual_input_required else "ready_for_generation",
            asset_type=asset_type,
            chapter_number=chapter_number,
            chapter_title=chapter_title,
            paragraph_index=index,
            placeholder=placeholder,
            nearby_context=context,
            prompt=build_prompt(
                asset_type,
                chapter_number,
                chapter_title,
                context,
                notes,
            ),
            manual_input_required=manual_input_required,
            notes=notes,
        )
        tasks.append(task)

    return tasks


def write_outputs(tasks: list[AssetTask]) -> tuple[Path, Path]:
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    json_path = OUTPUT_FOLDER / "asset_manifest.json"
    csv_path = OUTPUT_FOLDER / "asset_manifest.csv"

    records = [asdict(task) for task in tasks]

    json_path.write_text(
        json.dumps(records, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )

    with csv_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.DictWriter(file, fieldnames=records[0].keys())
        writer.writeheader()
        writer.writerows(records)

    return json_path, csv_path


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Create AI-ready asset tasks from handbook placeholders."
    )
    parser.add_argument(
        "source",
        nargs="?",
        type=Path,
        help="Optional source handbook DOCX. Defaults to handbook/source.",
    )
    args = parser.parse_args()

    try:
        source = args.source.resolve() if args.source else find_source()

        if not source.exists():
            raise FileNotFoundError(f"Source handbook not found: {source}")

        tasks = build_manifest(source)
    except (FileNotFoundError, ValueError) as error:
        print(f"Manifest generation failed: {error}")
        return 2

    if not tasks:
        print("No placeholders found. No asset manifest was created.")
        return 0

    json_path, csv_path = write_outputs(tasks)
    ready = sum(task.status == "ready_for_generation" for task in tasks)
    needs_input = len(tasks) - ready

    print("=" * 60)
    print("Vietnam Adventure Publisher")
    print("Asset Manifest")
    print("=" * 60)
    print(f"\nSource              : {source}")
    print(f"Asset tasks         : {len(tasks)}")
    print(f"Ready for generation: {ready}")
    print(f"Needs manual input  : {needs_input}")
    print(f"\nJSON: {json_path}")
    print(f"CSV : {csv_path}")

    return 0


if __name__ == "__main__":
    sys.exit(main())
