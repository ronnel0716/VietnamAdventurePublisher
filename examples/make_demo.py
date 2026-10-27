"""Generate synthetic input; no personal travel handbook is needed."""
from pathlib import Path
from docx import Document

path = Path(__file__).with_name('demo.docx')
if path.exists():
    raise SystemExit('demo.docx already exists; keep it or remove it explicitly before regenerating.')
doc = Document()
doc.add_heading('A Fictional Weekend Handbook', 0)
doc.add_paragraph('Synthetic portfolio fixture. All destinations and costs below are illustrative.')
doc.add_heading('Day 1 — Explore the city', 1)
doc.add_paragraph('Take a morning walk, visit a museum, then review the next day’s plan.')
doc.add_heading('Illustrative budget', 2)
table = doc.add_table(rows=1, cols=2)
table.rows[0].cells[0].text = 'Item'
table.rows[0].cells[1].text = 'Demo amount'
for name, amount in [('Transport', '₱250'), ('Meals', '₱500')]:
    cells = table.add_row().cells
    cells[0].text, cells[1].text = name, amount
doc.add_heading('Day 2 — Review and depart', 1)
doc.add_paragraph('Check your packing list and transport arrangements. This is not travel advice.')
doc.save(path)
print(path)
