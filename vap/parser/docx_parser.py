"""
docx_parser.py

Loads and parses the Vietnam Adventure Handbook.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from vap.models.handbook import Handbook
from vap.parser.chapter_detector import ChapterDetector


class DocxParser:
    """
    Reads a DOCX handbook and converts it into a Handbook object.
    """

    def __init__(self, filename: str | Path):

        self.filename = Path(filename)

        self.document = None
        self.handbook = Handbook()

    def load(self) -> Handbook:
        """
        Load the DOCX file and build the handbook model.
        """

        print(f"\nLoading: {self.filename.name}")

        self.document = Document(self.filename)

        self.handbook.title = self._get_title()

        self.handbook.paragraph_count = len(self.document.paragraphs)

        self.handbook.table_count = len(self.document.tables)

        detector = ChapterDetector()

        self.handbook.chapters = detector.detect(self.document)

        return self.handbook

    def _get_title(self) -> str:
        """
        Returns the first paragraph that uses the Title style.

        Falls back to the filename if no title is found.
        """

        for paragraph in self.document.paragraphs:

            if paragraph.style.name == "Title":

                text = paragraph.text.strip()

                if text:
                    return text

        return self.filename.stem

    @property
    def paragraphs(self):
        """
        Convenience accessor.
        """

        return self.document.paragraphs

    @property
    def tables(self):
        """
        Convenience accessor.
        """

        return self.document.tables