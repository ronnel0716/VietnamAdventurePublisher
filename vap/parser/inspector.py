from collections import Counter
from docx import Document


class DocumentInspector:

    def inspect(self, filename):

        doc = Document(filename)

        style_counter = Counter()

        print("=" * 70)
        print("DOCUMENT INSPECTOR")
        print("=" * 70)

        print()

        print(f"Paragraphs : {len(doc.paragraphs)}")
        print(f"Tables     : {len(doc.tables)}")

        print()

        for i, p in enumerate(doc.paragraphs, start=1):

            style = p.style.name

            style_counter[style] += 1

            text = p.text.strip()

            if len(text) > 80:
                text = text[:80] + "..."

            print(
                f"{i:04d} | "
                f"{style:<20} | "
                f"{text}"
            )

        print()

        print("=" * 70)
        print("STYLE SUMMARY")
        print("=" * 70)

        for style, count in sorted(style_counter.items()):
            print(f"{style:<30}{count}")