"""
document_reader.py

Reads a Microsoft Word document and converts it into a DocumentModel.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from vap.models.document import DocumentModel
from vap.models.paragraph import ParagraphModel


class DocumentReader:
    """
    Reads a DOCX file and converts it into a DocumentModel.
    """

    def __init__(self, filename: str | Path):

        self.filename = Path(filename)

    def read(self) -> DocumentModel:
        """
        Reads the DOCX file and returns a DocumentModel.
        """

        doc = Document(self.filename)

        document = DocumentModel()

        document.title = self._get_title(doc)
        document.paragraph_count = len(doc.paragraphs)
        document.table_count = len(doc.tables)

        for index, paragraph in enumerate(doc.paragraphs):

            document.add_paragraph(
                ParagraphModel(
                    index=index,
                    text=paragraph.text,
                    style=paragraph.style.name,
                )
            )

        return document

    @staticmethod
    def _get_title(doc) -> str:
        """
        Returns the first paragraph using the 'Title' style.

        Falls back to 'Untitled Document' if none exists.
        """

        for paragraph in doc.paragraphs:

            if paragraph.style.name == "Title":

                text = paragraph.text.strip()

                if text:
                    return text

        return "Untitled Document"