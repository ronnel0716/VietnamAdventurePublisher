"""
word_publisher.py

Publishes a semantic Handbook to a Microsoft Word document.

This is the initial implementation that writes the semantic
structure to a python-docx Document. Formatting and styling
will be expanded in later iterations.
"""

from docx import Document


class WordPublisher:
    """
    Publish a Handbook model to a DOCX document.
    """

    def publish(self, handbook, output_path=None):
        """
        Create a Word document from the Handbook model.

        Parameters
        ----------
        handbook
            Handbook model.

        output_path : str | None
            If supplied, saves the document.

        Returns
        -------
        Document
            python-docx Document instance.
        """

        doc = Document()

        title = getattr(handbook, "title", "")

        if title:
            doc.add_heading(title, level=0)

        for chapter in getattr(handbook, "chapters", []):

            chapter_title = getattr(chapter, "title", "Untitled Chapter")

            doc.add_heading(chapter_title, level=1)

            for component in getattr(chapter, "components", []):

                text = getattr(component, "text", "")

                if text:
                    doc.add_paragraph(text)

        if output_path:
            doc.save(output_path)

        return doc