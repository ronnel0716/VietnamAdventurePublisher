"""
reader.py

Reads a DOCX file and converts it into the application's
DocumentModel.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document

from vap.models.document import DocumentModel
from vap.models.paragraph import ParagraphModel


class DocumentReader:
    """
    Reads Microsoft Word documents and converts them into
    the internal DocumentModel.
    """

    def read(self, path: str | Path) -> DocumentModel:
        """
        Read a Microsoft Word document.

        Args:
            path:
                Path to the DOCX file.

        Returns:
            DocumentModel
        """

        path = Path(path)

        if not path.exists():
            raise FileNotFoundError(path)

        doc = Document(path)

        model = DocumentModel()

        # ------------------------------------------------------------------
        # Document metadata
        # ------------------------------------------------------------------

        if doc.paragraphs:
            model.title = doc.paragraphs[0].text.strip()

        model.table_count = len(doc.tables)

        # ------------------------------------------------------------------
        # Paragraphs
        # ------------------------------------------------------------------

        for index, paragraph in enumerate(doc.paragraphs):

            model.add_paragraph(
                ParagraphModel(
                    index=index,
                    text=paragraph.text,
                    style=paragraph.style.name,
                )
            )

        model.paragraph_count = len(model)

        return model